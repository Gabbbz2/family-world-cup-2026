from dotenv import load_dotenv
from pathlib import Path
ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import io
import re
import uuid
import secrets
import logging
import bcrypt
import jwt
from datetime import datetime, timezone, timedelta
from typing import List, Optional, Literal, Dict, Any
from collections import defaultdict

import openpyxl
from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, UploadFile, File, Form
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, EmailStr, Field

# ============== Config ==============
JWT_ALGORITHM = "HS256"
JWT_SECRET = os.environ["JWT_SECRET"]
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "gabriella.bengtsson2@gmail.com").lower()
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "WorldCup2026!")
DATA_DIR = ROOT_DIR / "data"

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

app = FastAPI(title="Family World Cup 2026")
api = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ============== Helpers ==============
def hash_password(pw: str) -> str:
    return bcrypt.hashpw(pw.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False

def now_utc() -> datetime:
    return datetime.now(timezone.utc)

def create_access_token(user_id: str, email: str) -> str:
    return jwt.encode(
        {"sub": user_id, "email": email,
         "exp": now_utc() + timedelta(days=7), "type": "access"},
        JWT_SECRET, algorithm=JWT_ALGORITHM,
    )

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(401, "Invalid token type")
        user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
        if not user:
            raise HTTPException(401, "User not found")
        if user.get("status") in ("deactivated", "banned", "deleted"):
            raise HTTPException(403, "Kontot är inte aktivt")
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")

async def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(403, "Admin access required")
    return user

def set_auth_cookie(response: Response, token: str):
    response.set_cookie(
        "access_token", token, httponly=True, secure=True,
        samesite="none", max_age=7 * 24 * 60 * 60, path="/",
    )

async def log_audit(admin: dict, action: str, details: dict = None):
    await db.audit_log.insert_one({
        "id": str(uuid.uuid4()),
        "admin_id": admin["id"],
        "admin_name": admin.get("name"),
        "admin_email": admin.get("email"),
        "action": action,
        "details": details or {},
        "timestamp": now_utc().isoformat(),
    })

# ============== Models ==============
class RegisterReq(BaseModel):
    name: str
    email: EmailStr
    password: str = Field(min_length=6)

class LoginReq(BaseModel):
    email: EmailStr
    password: str

class ForgotPwReq(BaseModel):
    email: EmailStr

class ResetPwReq(BaseModel):
    token: str
    password: str = Field(min_length=6)

class MatchCreate(BaseModel):
    stage: str = "group"
    round: Optional[str] = None
    group: Optional[str] = None
    home_team_id: Optional[str] = None
    away_team_id: Optional[str] = None
    home_placeholder: Optional[str] = None
    away_placeholder: Optional[str] = None
    kickoff: datetime
    tv_channel: Optional[str] = None
    match_number: Optional[int] = None

class MatchResultReq(BaseModel):
    home_score: int = Field(ge=0)
    away_score: int = Field(ge=0)

class PredictionReq(BaseModel):
    match_id: str
    home_score: int = Field(ge=0)
    away_score: int = Field(ge=0)

class TournamentPredictionReq(BaseModel):
    version: int = Field(ge=1, le=4)
    group_rankings: Dict[str, List[str]] = {}  # {"A": [1st_id, 2nd_id, 3rd_id, 4th_id]}
    advancing: List[str] = []
    r32: List[str] = []
    r16: List[str] = []
    qf: List[str] = []
    sf: List[str] = []
    third_place: List[str] = []
    finalists: List[str] = []
    champion: Optional[str] = None

class DisqualifyReq(BaseModel):
    team_id: str
    reason: str

class InviteReq(BaseModel):
    email: EmailStr

class StrategyReq(BaseModel):
    user_id: str
    points: int

class V1DeadlineReq(BaseModel):
    deadline: datetime

class DeadlineReq(BaseModel):
    deadline: datetime

class UserStatusReq(BaseModel):
    status: Literal["active", "deactivated", "banned"]
    reason: Optional[str] = None

class UserDeleteReq(BaseModel):
    confirmation: str  # must be "DELETE"
    reason: Optional[str] = None

class HallOfFameReq(BaseModel):
    year: int
    winner_name: str
    winner_points: int
    second_name: str
    second_points: int
    third_name: str
    third_points: int

# ============== Auth ==============
@api.post("/auth/register")
async def register(req: RegisterReq, response: Response):
    email = req.email.lower()
    existing = await db.users.find_one({"email": email})
    if existing:
        if existing.get("status") == "banned":
            raise HTTPException(403, "Den här e-postadressen är blockerad.")
        raise HTTPException(400, "E-postadressen är redan registrerad")
    invite = await db.invites.find_one({"email": email})
    if email != ADMIN_EMAIL and not invite:
        raise HTTPException(403, "Detta är en privat app. Din e-post finns inte på inbjudningslistan.")
    role = "admin" if email == ADMIN_EMAIL else "user"
    user = {
        "id": str(uuid.uuid4()),
        "name": req.name.strip(),
        "email": email,
        "password_hash": hash_password(req.password),
        "role": role,
        "status": "active",
        "created_at": now_utc().isoformat(),
        "last_login": now_utc().isoformat(),
        "live_points": 0,
        "strategy_points": 0,
    }
    await db.users.insert_one(user)
    if invite:
        await db.invites.delete_one({"email": email})
    token = create_access_token(user["id"], user["email"])
    set_auth_cookie(response, token)
    user.pop("password_hash", None)
    user.pop("_id", None)
    return {"user": user, "token": token}

@api.post("/auth/login")
async def login(req: LoginReq, response: Response):
    email = req.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(req.password, user["password_hash"]):
        raise HTTPException(401, "Felaktig e-post eller lösenord")
    status = user.get("status", "active")
    if status == "deactivated":
        raise HTTPException(403, "Ditt konto är avaktiverat. Kontakta administratören.")
    if status == "banned":
        raise HTTPException(403, "Ditt konto är blockerat. Kontakta administratören.")
    if status == "deleted":
        raise HTTPException(403, "Kontot är borttaget.")
    await db.users.update_one({"id": user["id"]}, {"$set": {"last_login": now_utc().isoformat()}})
    user["last_login"] = now_utc().isoformat()
    token = create_access_token(user["id"], user["email"])
    set_auth_cookie(response, token)
    user.pop("password_hash", None)
    user.pop("_id", None)
    return {"user": user, "token": token}

@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}

@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user

@api.post("/auth/forgot-password")
async def forgot_password(req: ForgotPwReq):
    email = req.email.lower()
    user = await db.users.find_one({"email": email})
    if not user:
        return {"ok": True}
    token = secrets.token_urlsafe(32)
    await db.password_reset_tokens.insert_one({
        "token": token, "user_id": user["id"],
        "expires_at": (now_utc() + timedelta(hours=1)).isoformat(), "used": False,
    })
    print(f"[PASSWORD RESET] {email} -> {token}")
    return {"ok": True, "reset_token": token}

@api.post("/auth/reset-password")
async def reset_password(req: ResetPwReq):
    doc = await db.password_reset_tokens.find_one({"token": req.token, "used": False})
    if not doc:
        raise HTTPException(400, "Invalid or used token")
    if datetime.fromisoformat(doc["expires_at"]) < now_utc():
        raise HTTPException(400, "Token expired")
    await db.users.update_one({"id": doc["user_id"]}, {"$set": {"password_hash": hash_password(req.password)}})
    await db.password_reset_tokens.update_one({"token": req.token}, {"$set": {"used": True}})
    return {"ok": True}

# ============== Teams ==============
@api.get("/teams")
async def list_teams():
    teams = await db.teams.find({}, {"_id": 0}).to_list(200)
    return sorted(teams, key=lambda t: (t.get("group", "Z"), t["team_name"]))

@api.get("/teams/groups")
async def teams_by_group():
    teams = await db.teams.find({}, {"_id": 0}).to_list(200)
    groups: Dict[str, list] = {}
    for t in teams:
        groups.setdefault(t["group"], []).append(t)
    for g in groups:
        groups[g].sort(key=lambda x: x["team_name"])
    return groups

# ============== Matches ==============
async def _attach_teams_to_matches(matches: list) -> list:
    ids = list({m["home_team_id"] for m in matches if m.get("home_team_id")} |
               {m["away_team_id"] for m in matches if m.get("away_team_id")})
    teams = await db.teams.find({"id": {"$in": ids}}, {"_id": 0}).to_list(200)
    tmap = {t["id"]: t for t in teams}
    for m in matches:
        m["home_team"] = tmap.get(m.get("home_team_id"))
        m["away_team"] = tmap.get(m.get("away_team_id"))
    return matches

@api.get("/matches")
async def list_matches():
    matches = await db.matches.find({}, {"_id": 0}).to_list(2000)
    matches = await _attach_teams_to_matches(matches)
    matches.sort(key=lambda m: (m.get("kickoff") or "", m.get("match_number") or 0))
    return matches

@api.post("/matches", dependencies=[Depends(require_admin)])
async def create_match(req: MatchCreate, admin: dict = Depends(require_admin)):
    m = {
        "id": str(uuid.uuid4()),
        "stage": req.stage,
        "round": req.round,
        "group": req.group,
        "home_team_id": req.home_team_id,
        "away_team_id": req.away_team_id,
        "home_placeholder": req.home_placeholder,
        "away_placeholder": req.away_placeholder,
        "kickoff": req.kickoff.isoformat() if isinstance(req.kickoff, datetime) else req.kickoff,
        "tv_channel": req.tv_channel,
        "match_number": req.match_number,
        "home_score": None, "away_score": None,
        "status": "upcoming",
    }
    await db.matches.insert_one(m)
    await log_audit(admin, "match_create", {"match_number": req.match_number})
    m.pop("_id", None)
    return m

@api.put("/matches/{match_id}/result")
async def set_result(match_id: str, req: MatchResultReq, admin: dict = Depends(require_admin)):
    m = await db.matches.find_one({"id": match_id}, {"_id": 0})
    if not m:
        raise HTTPException(404, "Match not found")
    prev = {"home_score": m.get("home_score"), "away_score": m.get("away_score")}
    await db.matches.update_one(
        {"id": match_id},
        {"$set": {"home_score": req.home_score, "away_score": req.away_score, "status": "finished"}},
    )
    await log_audit(admin, "match_result_set", {
        "match_id": match_id, "match_number": m.get("match_number"),
        "previous": prev, "new": {"home_score": req.home_score, "away_score": req.away_score},
    })
    await recompute_live_points()
    await progress_tournament()
    return {"ok": True}

@api.delete("/matches/{match_id}/result")
async def clear_result(match_id: str, admin: dict = Depends(require_admin)):
    m = await db.matches.find_one({"id": match_id}, {"_id": 0})
    if not m:
        raise HTTPException(404, "Match not found")
    await db.matches.update_one({"id": match_id}, {"$set": {"home_score": None, "away_score": None, "status": "upcoming"}})
    await log_audit(admin, "match_result_cleared", {"match_id": match_id, "match_number": m.get("match_number")})
    await recompute_live_points()
    await progress_tournament()
    return {"ok": True}

@api.delete("/matches/{match_id}", dependencies=[Depends(require_admin)])
async def delete_match(match_id: str):
    await db.matches.delete_one({"id": match_id})
    await db.match_predictions.delete_many({"match_id": match_id})
    return {"ok": True}

# ============== Predictions ==============
@api.post("/predictions")
async def submit_prediction(req: PredictionReq, user: dict = Depends(get_current_user)):
    match = await db.matches.find_one({"id": req.match_id}, {"_id": 0})
    if not match:
        raise HTTPException(404, "Match not found")
    if not match.get("home_team_id") or not match.get("away_team_id"):
        raise HTTPException(400, "Teams not yet determined")
    kickoff = datetime.fromisoformat(match["kickoff"]) if isinstance(match["kickoff"], str) else match["kickoff"]
    if kickoff.tzinfo is None:
        kickoff = kickoff.replace(tzinfo=timezone.utc)
    if now_utc() >= kickoff - timedelta(minutes=5):
        raise HTTPException(400, "Tippning stängd")
    doc = {
        "id": str(uuid.uuid4()), "user_id": user["id"], "match_id": req.match_id,
        "home_score": req.home_score, "away_score": req.away_score,
        "submitted_at": now_utc().isoformat(),
    }
    await db.match_predictions.update_one(
        {"user_id": user["id"], "match_id": req.match_id},
        {"$set": doc}, upsert=True,
    )
    return doc

@api.get("/predictions/me")
async def my_predictions(user: dict = Depends(get_current_user)):
    return await db.match_predictions.find({"user_id": user["id"]}, {"_id": 0}).to_list(2000)

@api.get("/predictions/match/{match_id}")
async def match_predictions_visible(match_id: str, user: dict = Depends(get_current_user)):
    match = await db.matches.find_one({"id": match_id}, {"_id": 0})
    if not match:
        raise HTTPException(404, "Match not found")
    kickoff = match.get("kickoff")
    if kickoff and user.get("role") != "admin":
        lock_time = kickoff - timedelta(minutes=5)
        if datetime.utcnow() < lock_time:
            return {"locked": True, "predictions": []}
    preds = await db.match_predictions.find({"match_id": match_id}, {"_id": 0}).to_list(500)
    users = await db.users.find({"id": {"$in": [p["user_id"] for p in preds]}}, {"_id": 0, "password_hash": 0}).to_list(500)
    umap = {u["id"]: u for u in users}
    for p in preds:
        p["user"] = umap.get(p["user_id"], {})
    return {"locked": False, "predictions": preds}

# ============== Tournament Predictions ==============
@api.post("/tournament-predictions")
async def submit_tp(req: TournamentPredictionReq, user: dict = Depends(get_current_user)):
    # Check version is not locked (deadline passed)
    dl_iso = await get_deadline_value(req.version)
    deadline_passed = False
    if dl_iso:
        dl = datetime.fromisoformat(dl_iso)
        if dl.tzinfo is None:
            dl = dl.replace(tzinfo=timezone.utc)
        deadline_passed = now_utc() > dl
    if deadline_passed:
        raise HTTPException(400, f"Version {req.version} är låst och kan inte ändras")

    # Validate group rankings: each group must have exactly 4 unique team IDs that all belong to that group
    teams_in_groups = {}
    teams = await db.teams.find({}, {"_id": 0, "id": 1, "group": 1}).to_list(200)
    for t in teams:
        teams_in_groups.setdefault(t["group"], set()).add(t["id"])
    if req.group_rankings:
        for g, ids in req.group_rankings.items():
            if g not in teams_in_groups:
                raise HTTPException(400, f"Okänd grupp: {g}")
            if len(ids) != 4:
                raise HTTPException(400, f"Grupp {g}: exakt 4 lag måste rankas (1:a–4:e)")
            if any(not x for x in ids):
                raise HTTPException(400, f"Grupp {g}: alla fyra platser måste fyllas i")
            if len(set(ids)) != 4:
                raise HTTPException(400, f"Grupp {g}: varje lag får bara väljas en gång")
            for tid in ids:
                if tid not in teams_in_groups[g]:
                    raise HTTPException(400, f"Grupp {g}: ett av lagen tillhör inte gruppen")

    doc = {
        "id": str(uuid.uuid4()), "user_id": user["id"], "version": req.version,
        "group_rankings": req.group_rankings,
        "advancing": req.advancing, "r32": req.r32, "r16": req.r16, "qf": req.qf,
        "sf": req.sf, "third_place": req.third_place, "finalists": req.finalists, "champion": req.champion,
        "submitted_at": now_utc().isoformat(),
        "is_late": False,
    }
    await db.tournament_predictions.update_one(
        {"user_id": user["id"], "version": req.version},
        {"$set": doc}, upsert=True,
    )
    return doc

@api.get("/tournament-predictions/me")
async def my_tp(user: dict = Depends(get_current_user)):
    return await db.tournament_predictions.find({"user_id": user["id"]}, {"_id": 0}).to_list(10)

@api.get("/tournament-predictions/visible")
async def visible_tp(user: dict = Depends(get_current_user)):
    """Return other users' tournament predictions for each version where the requesting user's own version is locked (deadline passed)."""
    out = {}
    for v in (1, 2, 3, 4):
        dl_iso = await get_deadline_value(v)
        if not dl_iso:
            continue
        dl = datetime.fromisoformat(dl_iso)
        if dl.tzinfo is None:
            dl = dl.replace(tzinfo=timezone.utc)
        if now_utc() <= dl and user.get("role") != "admin":
            continue
        docs = await db.tournament_predictions.find({"version": v}, {"_id": 0}).to_list(2000)
        users = await db.users.find({"id": {"$in": [d["user_id"] for d in docs]}}, {"_id": 0, "name": 1, "id": 1}).to_list(2000)
        umap = {u["id"]: u for u in users}
        for d in docs:
            d["user_name"] = umap.get(d["user_id"], {}).get("name", "Spelare")
        out[str(v)] = docs
    return out

@api.get("/tournament-predictions/all", dependencies=[Depends(require_admin)])
async def all_tp():
    return await db.tournament_predictions.find({}, {"_id": 0}).to_list(2000)

# ============== Config (deadlines) ==============
DEADLINE_KEYS = {1: "v1_deadline", 2: "v2_deadline", 3: "v3_deadline", 4: "v4_deadline"}

async def get_deadline_value(version: int) -> Optional[str]:
    key = DEADLINE_KEYS.get(version)
    if not key:
        return None
    cfg = await db.config.find_one({"key": key}, {"_id": 0}) or {}
    return cfg.get("value")

@api.get("/config/v1-deadline")
async def get_v1_deadline():
    return {"deadline": await get_deadline_value(1)}

@api.post("/config/v1-deadline", dependencies=[Depends(require_admin)])
async def set_v1_deadline(req: V1DeadlineReq, admin: dict = Depends(require_admin)):
    dl = req.deadline.isoformat() if isinstance(req.deadline, datetime) else req.deadline
    prev = await get_deadline_value(1)
    await db.config.update_one({"key": "v1_deadline"}, {"$set": {"key": "v1_deadline", "value": dl}}, upsert=True)
    await log_audit(admin, "deadline_changed", {"version": 1, "old_value": prev, "new_value": dl})
    return {"ok": True}

@api.get("/config/deadlines")
async def get_all_deadlines():
    return {str(v): await get_deadline_value(v) for v in (1, 2, 3, 4)}

@api.post("/config/v{version}-deadline", dependencies=[Depends(require_admin)])
async def set_version_deadline(version: int, req: DeadlineReq, admin: dict = Depends(require_admin)):
    if version not in DEADLINE_KEYS:
        raise HTTPException(400, "Ogiltig version")
    key = DEADLINE_KEYS[version]
    dl = req.deadline.isoformat() if isinstance(req.deadline, datetime) else req.deadline
    prev = await get_deadline_value(version)
    await db.config.update_one({"key": key}, {"$set": {"key": key, "value": dl}}, upsert=True)
    await log_audit(admin, "deadline_changed", {"version": version, "old_value": prev, "new_value": dl})
    return {"ok": True}

# ============== Strategy Versions (new system) ==============
# A series of tipping rounds, each opening just before the relevant stage.
# pre_tournament keeps the legacy V1 shape (group_rankings + bracket lists + champion).
# All other versions store {picks: {match_id: team_id}} — winner-only, no scores.
VERSION_DEFS = [
    {"id": "pre_tournament", "label": "Första Tipset", "order": 1, "round_filter": None, "stage_filter": "Group Stage", "points_per_correct": None},
    {"id": "r32",            "label": "Tips Sextondelsfinal", "order": 2, "round_filter": "sextondelsfinal", "stage_filter": "Knockout", "points_per_correct": 8},
    {"id": "r16",            "label": "Tips Åttondelsfinal",  "order": 3, "round_filter": "åttondelsfinal",  "stage_filter": "Knockout", "points_per_correct": 6},
    {"id": "qf",             "label": "Tips Kvartsfinal",     "order": 4, "round_filter": "kvartsfinal",     "stage_filter": "Knockout", "points_per_correct": 5},
    {"id": "sf",             "label": "Tips Semifinal",       "order": 5, "round_filter": "semifinal",       "stage_filter": "Knockout", "points_per_correct": 4},
    {"id": "third_place",    "label": "Tips Bronsmatch",      "order": 6, "round_filter": "bronsmatch",      "stage_filter": "Knockout", "points_per_correct": 3},
    {"id": "final",          "label": "Tips Final",           "order": 7, "round_filter": "final",           "stage_filter": "Knockout", "points_per_correct": 5},
]
VERSION_BY_ID = {v["id"]: v for v in VERSION_DEFS}
VALID_VERSION_IDS = [v["id"] for v in VERSION_DEFS]

class StrategySubmitReq(BaseModel):
    # For pre_tournament:
    group_rankings: Dict[str, List[str]] = {}
    advancing: List[str] = []
    r32: List[str] = []
    r16: List[str] = []
    qf: List[str] = []
    sf: List[str] = []
    third_place: List[str] = []  # bronsmatch
    finalists: List[str] = []
    champion: Optional[str] = None
    # For knockout versions:
    picks: Dict[str, str] = {}  # match_id -> team_id (winner)

class StrategyOverrideReq(BaseModel):
    status: Literal["auto", "open", "closed"] = "auto"
    custom_deadline: Optional[datetime] = None

async def matches_for_version(version_id: str) -> List[dict]:
    vdef = VERSION_BY_ID.get(version_id)
    if not vdef:
        return []
    if version_id == "pre_tournament":
        ms = await db.matches.find({"stage": vdef["stage_filter"]}, {"_id": 0}).to_list(2000)
    else:
        rf = vdef["round_filter"]
        ms = await db.matches.find(
            {"stage": "Knockout", "round": {"$regex": f"^{rf}", "$options": "i"}},
            {"_id": 0},
        ).to_list(500)
    ms.sort(key=lambda m: (m.get("kickoff") or "", m.get("match_number") or 0))
    return ms

async def compute_default_deadline(version_id: str) -> Optional[str]:
    ms = await matches_for_version(version_id)
    if not ms:
        return None
    kickoff_iso = ms[0]["kickoff"]
    kickoff = datetime.fromisoformat(kickoff_iso)
    if kickoff.tzinfo is None:
        kickoff = kickoff.replace(tzinfo=timezone.utc)
    return (kickoff - timedelta(minutes=5)).isoformat()

async def get_version_state(version_id: str) -> dict:
    """Return {deadline, locked, override} for a version."""
    default_dl = await compute_default_deadline(version_id)
    override = await db.config.find_one({"key": f"strategy_override_{version_id}"}, {"_id": 0}) or {}
    status = override.get("status", "auto")
    custom = override.get("custom_deadline")
    effective_dl = custom or default_dl
    if version_id == "pre_tournament" and status != "closed":
        locked = False
    elif status == "open":
        locked = False
    elif status == "closed":
        locked = True
    else:
        if effective_dl:
            dl = datetime.fromisoformat(effective_dl)
            if dl.tzinfo is None:
                dl = dl.replace(tzinfo=timezone.utc)
            locked = now_utc() >= dl
        else:
            locked = False
    return {
        "default_deadline": default_dl,
        "deadline": effective_dl,
        "override_status": status,
        "locked": locked,
    }

@api.get("/strategy/versions")
async def list_strategy_versions(user: dict = Depends(get_current_user)):
    out = []
    my_subs = {d["version"]: d for d in await db.tournament_predictions.find({"user_id": user["id"]}, {"_id": 0}).to_list(20)}
    for v in VERSION_DEFS:
        state = await get_version_state(v["id"])
        ms = await matches_for_version(v["id"])
        out.append({
            **v,
            **state,
            "match_count": len(ms),
            "my_submitted": v["id"] in my_subs,
            "my_submission": my_subs.get(v["id"]),
        })
    return out

@api.get("/strategy/version/{version_id}/matches")
async def matches_for_strategy_version(version_id: str):
    if version_id not in VALID_VERSION_IDS:
        raise HTTPException(404, "Okänd version")
    ms = await matches_for_version(version_id)
    return await _attach_teams_to_matches(ms)

@api.post("/strategy/version/{version_id}")
async def submit_strategy_version(version_id: str, req: StrategySubmitReq, user: dict = Depends(get_current_user)):
    if version_id not in VALID_VERSION_IDS:
        raise HTTPException(404, "Okänd version")
    state = await get_version_state(version_id)
    if state["locked"]:
        raise HTTPException(400, "Tippning stängd")

    doc = {
        "id": str(uuid.uuid4()),
        "user_id": user["id"],
        "version": version_id,
        "submitted_at": now_utc().isoformat(),
    }

    if version_id == "pre_tournament":
        # Reuse old pre-tournament shape with full validation
        teams_in_groups: Dict[str, set] = {}
        teams = await db.teams.find({}, {"_id": 0, "id": 1, "group": 1}).to_list(200)
        for t in teams:
            teams_in_groups.setdefault(t["group"], set()).add(t["id"])
        if req.group_rankings:
            for g, ids in req.group_rankings.items():
                if g not in teams_in_groups:
                    raise HTTPException(400, f"Okänd grupp: {g}")
                if len(ids) != 4:
                    raise HTTPException(400, f"Grupp {g}: exakt 4 lag måste rankas (1:a–4:e)")
                if any(not x for x in ids):
                    raise HTTPException(400, f"Grupp {g}: alla fyra platser måste fyllas i")
                if len(set(ids)) != 4:
                    raise HTTPException(400, f"Grupp {g}: varje lag får bara väljas en gång")
                for tid in ids:
                    if tid not in teams_in_groups[g]:
                        raise HTTPException(400, f"Grupp {g}: ett av lagen tillhör inte gruppen")
        doc.update({
            "group_rankings": req.group_rankings,
            "advancing": req.advancing, "r32": req.r32, "r16": req.r16, "qf": req.qf,
            "sf": req.sf, "third_place": req.third_place, "finalists": req.finalists, "champion": req.champion,
        })
    else:
        # Knockout winner picks: validate each match belongs to this version's stage
        valid_match_ids = {m["id"] for m in await matches_for_version(version_id)}
        match_team_map: Dict[str, set] = {}
        for m in await matches_for_version(version_id):
            allowed = {m.get("home_team_id"), m.get("away_team_id")} - {None}
            match_team_map[m["id"]] = allowed
        cleaned: Dict[str, str] = {}
        for mid, tid in (req.picks or {}).items():
            if mid not in valid_match_ids:
                raise HTTPException(400, f"Match {mid} ingår inte i denna version")
            if not tid:
                continue
            allowed = match_team_map.get(mid, set())
            if allowed and tid not in allowed:
                raise HTTPException(400, "Vinnaren måste vara ett av matchens lag")
            cleaned[mid] = tid
        doc["picks"] = cleaned

    await db.tournament_predictions.update_one(
        {"user_id": user["id"], "version": version_id},
        {"$set": doc}, upsert=True,
    )
    return doc

@api.get("/strategy/version/{version_id}/predictions")
async def visible_strategy_predictions(version_id: str, user: dict = Depends(get_current_user)):
    if version_id not in VALID_VERSION_IDS:
        raise HTTPException(404, "Okänd version")
    state = await get_version_state(version_id)
    if not state["locked"] and user.get("role") != "admin":
        return {"locked": True, "predictions": []}
    docs = await db.tournament_predictions.find({"version": version_id}, {"_id": 0}).to_list(2000)
    users = await db.users.find({"id": {"$in": [d["user_id"] for d in docs]}}, {"_id": 0, "name": 1, "id": 1}).to_list(2000)
    umap = {u["id"]: u for u in users}
    for d in docs:
        d["user_name"] = umap.get(d["user_id"], {}).get("name", "Spelare")
    return {"locked": False, "predictions": docs}

@api.get("/strategy/remaining-teams/{after_stage}")
async def get_remaining_teams_after_stage(after_stage: str):
    """
    Returns teams still remaining in tournament after a given stage completes.
    after_stage can be: "group", "r32", "r16", "qf", "sf", "third_place"
    Returns: {stage: [team_id, ...]} for all future stages
    """
    valid_stages = ["group", "r32", "r16", "qf", "sf", "third_place"]
    if after_stage not in valid_stages:
        raise HTTPException(400, f"Okänd omgång: {after_stage}")
    
    matches = await db.matches.find({}, {"_id": 0}).to_list(2000)
    
    # Build set of teams for each stage based on completed matches
    stage_teams = {
        "r32": set(), "r16": set(), "qf": set(), 
        "sf": set(), "third_place": set(), "final": set(), "champion": set()
    }
    
    for m in matches:
        if m.get("status") != "finished":
            continue
        stage_l = (m.get("round") or m.get("stage") or "").lower()
        for tid_key in ("home_team_id", "away_team_id"):
            tid = m.get(tid_key)
            if not tid:
                continue
            # Determine winner
            hs = m.get("home_score")
            as_ = m.get("away_score")
            if hs is None or as_ is None:
                continue
            winner_id = None
            if hs > as_:
                winner_id = m.get("home_team_id")
            elif as_ > hs:
                winner_id = m.get("away_team_id")
            
            # Mark both teams as participants, but only winner advances
            if "sextondelsfinal" in stage_l or m.get("stage") == "r32":
                stage_teams["r32"].add(tid)
                if winner_id:
                    stage_teams["r16"].add(winner_id)
            elif "åttondelsfinal" in stage_l or m.get("stage") == "r16":
                stage_teams["r16"].add(tid)
                if winner_id:
                    stage_teams["qf"].add(winner_id)
            elif "kvartsfinal" in stage_l or m.get("stage") == "qf":
                stage_teams["qf"].add(tid)
                if winner_id:
                    stage_teams["sf"].add(winner_id)
            elif "semifinal" in stage_l or m.get("stage") == "sf":
                stage_teams["sf"].add(tid)
                if winner_id:
                    stage_teams["third_place"].add(winner_id)
                    stage_teams["final"].add(winner_id)
            elif "bronsmatch" in stage_l or m.get("round") == "bronsmatch":
                stage_teams["third_place"].add(tid)
            elif stage_l == "final" or m.get("stage") == "final":
                stage_teams["final"].add(tid)
                stage_teams["champion"].add(tid)
    
    # Return remaining stages after the given stage
    remaining_stages = []
    include_from = False
    stage_order = ["group", "r32", "r16", "qf", "sf", "third_place", "final", "champion"]
    
    for stage in stage_order:
        if stage == after_stage:
            include_from = True
            continue
        if include_from and stage in stage_teams:
            remaining_stages.append(stage)
    
    result = {}
    for stage in remaining_stages:
        result[stage] = list(stage_teams.get(stage, []))
    
    return result

@api.post("/admin/strategy/version/{version_id}/override")
async def admin_strategy_override(version_id: str, req: StrategyOverrideReq, admin: dict = Depends(require_admin)):
    if version_id not in VALID_VERSION_IDS:
        raise HTTPException(404, "Okänd version")
    key = f"strategy_override_{version_id}"
    prev = await db.config.find_one({"key": key}, {"_id": 0}) or {}
    custom = req.custom_deadline.isoformat() if isinstance(req.custom_deadline, datetime) else req.custom_deadline
    await db.config.update_one(
        {"key": key},
        {"$set": {"key": key, "status": req.status, "custom_deadline": custom}},
        upsert=True,
    )
    await log_audit(admin, "strategy_version_override", {
        "version_id": version_id,
        "old_value": {"status": prev.get("status"), "custom_deadline": prev.get("custom_deadline")},
        "new_value": {"status": req.status, "custom_deadline": custom},
    })
    return {"ok": True}

@api.post("/admin/strategy/recompute")
async def admin_strategy_recompute(admin: dict = Depends(require_admin)):
    await recompute_strategy_points()
    await log_audit(admin, "strategy_recompute")
    return {"ok": True}

# ============== Scoring ==============
def score_match(pred_h, pred_a, actual_h, actual_a) -> int:
    if actual_h is None or actual_a is None:
        return 0
    pts = 0
    pw = (pred_h > pred_a) - (pred_h < pred_a)
    aw = (actual_h > actual_a) - (actual_h < actual_a)
    if pw == aw:
        pts += 3
    if (pred_h - pred_a) == (actual_h - actual_a):
        pts += 2
    if pred_h == actual_h and pred_a == actual_a:
        pts += 5
    return pts

async def recompute_live_points():
    matches = await db.matches.find({"status": "finished"}, {"_id": 0}).to_list(2000)
    mmap = {m["id"]: m for m in matches}
    preds = await db.match_predictions.find({}, {"_id": 0}).to_list(20000)
    user_points: Dict[str, int] = {}
    for p in preds:
        m = mmap.get(p["match_id"])
        if not m:
            continue
        pts = score_match(p["home_score"], p["away_score"], m.get("home_score"), m.get("away_score"))
        user_points[p["user_id"]] = user_points.get(p["user_id"], 0) + pts
    all_users = await db.users.find({}, {"_id": 0, "id": 1}).to_list(2000)
    for u in all_users:
        await db.users.update_one({"id": u["id"]}, {"$set": {"live_points": user_points.get(u["id"], 0)}})

# Strategy scoring constants
STRATEGY_POINTS = {
    "group_winner": 5, "advancing": 3, "r32": 4, "r16": 6,
    "qf": 8, "sf": 12, "third_place": 15, "finalist": 20, "champion": 30,
}
VERSION_MULT = {1: 1.0, 2: 0.75, 3: 0.5, 4: 0.25}

async def recompute_strategy_points():
    """Compute strategy points for each user based on tournament predictions vs actual results so far."""
    # Determine actual outcomes from match data + standings
    standings = await compute_all_standings()

    # Compute completed groups and confirmed advancers to R32
    actual_group_winners: Dict[str, str] = {}
    actual_advancing: set = set()
    group_complete: Dict[str, bool] = {}
    for g, rows in standings.items():
        # group is considered complete only when every team has played 3 matches
        group_complete[g] = len(rows) >= 4 and all(r.get("played", 0) >= 3 for r in rows)
        if group_complete[g] and rows:
            actual_group_winners[g] = rows[0]["team_id"]
            # top two from complete groups advance
            for r in rows[:2]:
                actual_advancing.add(r["team_id"])
    # Best 8 third-placed teams only from complete groups
    thirds = [(g, rows[2]) for g, rows in standings.items() if group_complete.get(g) and len(rows) >= 3]
    thirds.sort(key=lambda x: (-x[1]["points"], -x[1]["goal_diff"], -x[1]["goals_for"]))
    for g, r in thirds[:8]:
        actual_advancing.add(r["team_id"])

    # Gather finished matches and compute confirmed participants per knockout stage
    matches = await db.matches.find({}, {"_id": 0}).to_list(2000)
    confirmed_r32 = set(actual_advancing)  # confirmed by completed groups
    confirmed_r16 = set()
    confirmed_qf = set()
    confirmed_sf = set()
    confirmed_finalists = set()
    bronze_winner_id: Optional[str] = None
    champion_id: Optional[str] = None

    # Map of finished knockout winners by match stage to derive confirmed progressions
    for m in matches:
        if m.get("status") != "finished":
            continue
        stage_l = (m.get("round") or m.get("stage") or "").lower()
        hs, as_ = m.get("home_score"), m.get("away_score")
        if hs is None or as_ is None:
            continue
        winner = m.get("home_team_id") if hs > as_ else (m.get("away_team_id") if as_ > hs else None)
        if not winner:
            continue
        # Finished R32 matches produce confirmed R16 participants
        if "sextondelsfinal" in stage_l or m.get("stage") == "r32":
            confirmed_r16.add(winner)
        # Finished R16 -> produce confirmed QF
        if "åttondelsfinal" in stage_l or m.get("stage") == "r16":
            confirmed_qf.add(winner)
        # Finished QF -> produce confirmed SF
        if "kvartsfinal" in stage_l or m.get("stage") == "qf":
            confirmed_sf.add(winner)
        # Finished SF -> finalists
        if "semifinal" in stage_l or m.get("stage") == "sf":
            confirmed_finalists.add(winner)
        # Bronze match winner
        if "bronsmatch" in stage_l or m.get("round") == "bronsmatch":
            bronze_winner_id = winner
        # Final winner -> champion
        if stage_l == "final" or m.get("stage") == "final":
            champion_id = winner

    # Pre-compute actual winner per finished knockout match (by match_id) for per-match picks
    actual_winner_by_match: Dict[str, str] = {}
    for m in matches:
        if m.get("stage") == "Knockout" and m.get("status") == "finished":
            hs, as_ = m.get("home_score"), m.get("away_score")
            if hs is None or as_ is None:
                continue
            if hs > as_:
                actual_winner_by_match[m["id"]] = m.get("home_team_id")
            elif as_ > hs:
                actual_winner_by_match[m["id"]] = m.get("away_team_id")

    # Score each user's tournament predictions
    tps = await db.tournament_predictions.find({}, {"_id": 0}).to_list(2000)
    user_strategy: Dict[str, float] = defaultdict(float)
    for tp in tps:
        version_id = tp.get("version", "pre_tournament")
        if isinstance(version_id, int) or (isinstance(version_id, str) and version_id.isdigit()):
            version_id = "pre_tournament"
        pts = 0
        if version_id == "pre_tournament":
            # Group rankings: 1st correct = +5 (only if group complete and winner known)
            for g, ranking in (tp.get("group_rankings") or {}).items():
                if ranking and len(ranking) >= 1 and group_complete.get(g) and ranking[0] == actual_group_winners.get(g):
                    pts += STRATEGY_POINTS["group_winner"]
                # Advancement points: award for any predicted team that actually advanced to R32
                for tid in (ranking or []):
                    if tid and tid in confirmed_r32:
                        pts += STRATEGY_POINTS["advancing"]

            # Bracket lists: only award when the stage progression is confirmed
            for tid in (tp.get("r32") or []):
                if tid in confirmed_r32:
                    pts += STRATEGY_POINTS["r32"]
            for tid in (tp.get("r16") or []):
                if tid in confirmed_r16:
                    pts += STRATEGY_POINTS["r16"]
            for tid in (tp.get("qf") or []):
                if tid in confirmed_qf:
                    pts += STRATEGY_POINTS["qf"]
            for tid in (tp.get("sf") or []):
                if tid in confirmed_sf:
                    pts += STRATEGY_POINTS["sf"]
            # Bronze match: award only if team won the bronze match
            for tid in (tp.get("third_place") or []):
                if bronze_winner_id and tid == bronze_winner_id:
                    pts += STRATEGY_POINTS["third_place"]
            # Finalists: award only when finalists are known (confirmed_finalists)
            for tid in (tp.get("finalists") or []):
                if tid in confirmed_finalists:
                    pts += STRATEGY_POINTS["finalist"]
            # Champion: award only when final completed and champion known
            if tp.get("champion") and champion_id and tp["champion"] == champion_id:
                pts += STRATEGY_POINTS["champion"]
        else:
            # Knockout-stage winner picks: only count finished matches
            vdef = VERSION_BY_ID.get(version_id)
            ppc = (vdef or {}).get("points_per_correct") or 0
            for mid, picked_tid in (tp.get("picks") or {}).items():
                actual = actual_winner_by_match.get(mid)
                if actual and picked_tid == actual:
                    pts += ppc
        user_strategy[tp["user_id"]] += pts

    all_users = await db.users.find({}, {"id": 1, "_id": 0}).to_list(2000)
    for u in all_users:
        await db.users.update_one({"id": u["id"]}, {"$set": {"strategy_points": int(round(user_strategy.get(u["id"], 0)))}})

# ============== Standings ==============
async def compute_all_standings() -> Dict[str, list]:
    """Return {group_letter: [ {team_id, team_name, country_code, played, won, drawn, lost, goals_for, goals_against, goal_diff, points}, ... ]}"""
    teams = await db.teams.find({}, {"_id": 0}).to_list(200)
    matches = await db.matches.find({"stage": {"$in": ["group", "Group Stage"]}, "status": "finished"}, {"_id": 0}).to_list(2000)
    by_group: Dict[str, Dict[str, dict]] = defaultdict(dict)
    for t in teams:
        g = t.get("group")
        if not g or t.get("disqualified"):
            continue
        by_group[g][t["id"]] = {
            "team_id": t["id"], "team_name": t["team_name"], "country_code": t["country_code"],
            "played": 0, "won": 0, "drawn": 0, "lost": 0,
            "goals_for": 0, "goals_against": 0, "goal_diff": 0, "points": 0,
        }
    for m in matches:
        g = m.get("group")
        if g is None:
            continue
        h, a = m.get("home_team_id"), m.get("away_team_id")
        if not h or not a:
            continue
        hs, as_ = m.get("home_score"), m.get("away_score")
        if hs is None or as_ is None:
            continue
        hrow = by_group[g].get(h)
        arow = by_group[g].get(a)
        if not hrow or not arow:
            continue
        hrow["played"] += 1
        arow["played"] += 1
        hrow["goals_for"] += hs
        hrow["goals_against"] += as_
        arow["goals_for"] += as_
        arow["goals_against"] += hs
        if hs > as_:
            hrow["won"] += 1
            arow["lost"] += 1
            hrow["points"] += 3
        elif as_ > hs:
            arow["won"] += 1
            hrow["lost"] += 1
            arow["points"] += 3
        else:
            hrow["drawn"] += 1
            arow["drawn"] += 1
            hrow["points"] += 1
            arow["points"] += 1
    out: Dict[str, list] = {}
    for g, teams_map in by_group.items():
        rows = list(teams_map.values())
        for r in rows:
            r["goal_diff"] = r["goals_for"] - r["goals_against"]
        rows.sort(key=lambda r: (-r["points"], -r["goal_diff"], -r["goals_for"], r["team_name"]))
        out[g] = rows
    return out

@api.get("/standings")
async def get_standings():
    return await compute_all_standings()

# ============== Tournament Progression ==============
PLACEHOLDER_GROUP_RE = re.compile(r"^([1234])([A-L])$")           # 1A, 2B, 3C
PLACEHOLDER_THIRD_RE = re.compile(r"^3([A-L]{2,})$")              # 3ABCDF (best third among listed groups)
PLACEHOLDER_W_RE = re.compile(r"^W(\d+)$")                        # W73
PLACEHOLDER_RU_RE = re.compile(r"^RU(\d+)$")                      # RU101

async def resolve_placeholder(label: Optional[str], standings: Dict[str, list], match_winners: Dict[int, str], match_losers: Dict[int, str], best_thirds: List[dict]) -> Optional[str]:
    if not label:
        return None
    m = PLACEHOLDER_GROUP_RE.match(label)
    if m:
        pos = int(m.group(1))
        g = m.group(2)
        rows = standings.get(g, [])
        if len(rows) >= pos:
            return rows[pos - 1]["team_id"]
        return None
    m = PLACEHOLDER_THIRD_RE.match(label)
    if m:
        groups = list(m.group(1))
        # Pick best 3rd among listed groups that is in qualified best-thirds
        candidates = [t for t in best_thirds if t["group"] in groups]
        if candidates:
            return candidates[0]["team_id"]
        return None
    m = PLACEHOLDER_W_RE.match(label)
    if m:
        return match_winners.get(int(m.group(1)))
    m = PLACEHOLDER_RU_RE.match(label)
    if m:
        return match_losers.get(int(m.group(1)))
    return None

async def progress_tournament():
    """Replace placeholders with real teams where possible."""
    # Group standings
    standings = await compute_all_standings()
    # Only consider a group complete when all 6 group matches finished (each team played 3)
    group_complete = {g: all(r["played"] >= 3 for r in rows) for g, rows in standings.items()}
    # Best thirds across all complete groups
    thirds_pool = []
    for g, rows in standings.items():
        if group_complete.get(g) and len(rows) >= 3:
            r = rows[2]
            thirds_pool.append({**r, "group": g})
    thirds_pool.sort(key=lambda r: (-r["points"], -r["goal_diff"], -r["goals_for"]))
    best_thirds = thirds_pool[:8]

    # Match winners/losers by match_number
    finished = await db.matches.find({"status": "finished"}, {"_id": 0}).to_list(2000)
    winners: Dict[int, str] = {}
    losers: Dict[int, str] = {}
    for m in finished:
        n = m.get("match_number")
        if n is None:
            continue
        hs, as_ = m.get("home_score"), m.get("away_score")
        if hs is None or as_ is None:
            continue
        if hs > as_:
            winners[n] = m.get("home_team_id")
            losers[n] = m.get("away_team_id")
        elif as_ > hs:
            winners[n] = m.get("away_team_id")
            losers[n] = m.get("home_team_id")

    # Iterate matches and resolve placeholders
    all_matches = await db.matches.find({}, {"_id": 0}).to_list(2000)
    updates = 0
    for m in all_matches:
        update = {}
        if not m.get("home_team_id") and m.get("home_placeholder"):
            # Only resolve group placeholders when group complete; W/RU when match finished
            label = m["home_placeholder"]
            mg = PLACEHOLDER_GROUP_RE.match(label) or PLACEHOLDER_THIRD_RE.match(label)
            if mg and isinstance(mg, re.Match) and PLACEHOLDER_GROUP_RE.match(label):
                if not group_complete.get(PLACEHOLDER_GROUP_RE.match(label).group(2)):
                    pass
                else:
                    tid = await resolve_placeholder(label, standings, winners, losers, best_thirds)
                    if tid:
                        update["home_team_id"] = tid
            elif PLACEHOLDER_THIRD_RE.match(label):
                # Need all listed groups complete
                groups = list(PLACEHOLDER_THIRD_RE.match(label).group(1))
                if all(group_complete.get(g) for g in groups):
                    tid = await resolve_placeholder(label, standings, winners, losers, best_thirds)
                    if tid:
                        update["home_team_id"] = tid
            else:
                tid = await resolve_placeholder(label, standings, winners, losers, best_thirds)
                if tid:
                    update["home_team_id"] = tid
        if not m.get("away_team_id") and m.get("away_placeholder"):
            label = m["away_placeholder"]
            if PLACEHOLDER_GROUP_RE.match(label):
                if group_complete.get(PLACEHOLDER_GROUP_RE.match(label).group(2)):
                    tid = await resolve_placeholder(label, standings, winners, losers, best_thirds)
                    if tid:
                        update["away_team_id"] = tid
            elif PLACEHOLDER_THIRD_RE.match(label):
                groups = list(PLACEHOLDER_THIRD_RE.match(label).group(1))
                if all(group_complete.get(g) for g in groups):
                    tid = await resolve_placeholder(label, standings, winners, losers, best_thirds)
                    if tid:
                        update["away_team_id"] = tid
            else:
                tid = await resolve_placeholder(label, standings, winners, losers, best_thirds)
                if tid:
                    update["away_team_id"] = tid
        if update:
            await db.matches.update_one({"id": m["id"]}, {"$set": update})
            updates += 1
    if updates:
        logger.info(f"Progression: updated {updates} matches with resolved placeholders")
    # Recompute strategy points after progression
    await recompute_strategy_points()

@api.post("/admin/progress", dependencies=[Depends(require_admin)])
async def admin_progress():
    await progress_tournament()
    return {"ok": True}

# ============== Leaderboard ==============
@api.get("/leaderboard")
async def leaderboard():
    # Only active users on leaderboard
    users = await db.users.find(
        {"$or": [{"status": "active"}, {"status": {"$exists": False}}]},
        {"_id": 0, "password_hash": 0},
    ).to_list(1000)
    rows = []
    for u in users:
        live = u.get("live_points", 0) or 0
        strat = u.get("strategy_points", 0) or 0
        rows.append({
            "user_id": u["id"], "name": u["name"], "email": u["email"],
            "live_points": live, "strategy_points": strat,
            "total_points": live + strat, "role": u.get("role", "user"),
        })
    rows.sort(key=lambda r: r["total_points"], reverse=True)
    rank = 1
    prev_points = None
    for i, r in enumerate(rows):
        if i > 0 and r["total_points"] != prev_points:
            rank += 1
        r["rank"] = rank
        prev_points = r["total_points"]
    return rows

# ============== Admin: Users & Invites ==============
@api.get("/admin/users", dependencies=[Depends(require_admin)])
async def admin_users():
    users = await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(1000)
    # Ensure status defaults to active for legacy docs
    for u in users:
        u.setdefault("status", "active")
    return users

@api.put("/admin/users/{user_id}/role")
async def set_user_role(user_id: str, role: dict, admin: dict = Depends(require_admin)):
    new_role = role.get("role")
    if new_role not in ("admin", "user"):
        raise HTTPException(400, "Ogiltig roll")
    u = await db.users.find_one({"id": user_id}, {"_id": 0, "role": 1, "name": 1, "email": 1})
    if not u:
        raise HTTPException(404, "Användaren hittades inte")
    await db.users.update_one({"id": user_id}, {"$set": {"role": new_role}})
    await log_audit(admin, "role_changed", {
        "user_id": user_id, "user_email": u.get("email"),
        "old_value": u.get("role"), "new_value": new_role,
    })
    return {"ok": True}

@api.put("/admin/users/{user_id}/status")
async def set_user_status(user_id: str, req: UserStatusReq, admin: dict = Depends(require_admin)):
    u = await db.users.find_one({"id": user_id}, {"_id": 0})
    if not u:
        raise HTTPException(404, "Användaren hittades inte")
    if req.status == "banned" and not (req.reason or "").strip():
        raise HTTPException(400, "Anledning krävs vid blockering")
    updates = {"status": req.status}
    if req.status == "banned":
        updates["ban_reason"] = req.reason
        updates["banned_at"] = now_utc().isoformat()
        updates["banned_by"] = admin.get("name")
    elif req.status == "active":
        updates["ban_reason"] = None
        updates["banned_at"] = None
        updates["banned_by"] = None
    await db.users.update_one({"id": user_id}, {"$set": updates})
    action_map = {
        "active": "user_activated", "deactivated": "user_deactivated", "banned": "user_banned",
    }
    await log_audit(admin, action_map[req.status], {
        "user_id": user_id, "user_email": u.get("email"),
        "old_value": u.get("status", "active"), "new_value": req.status,
        "reason": req.reason,
    })
    return {"ok": True}

@api.post("/admin/users/{user_id}/unban")
async def unban_user(user_id: str, admin: dict = Depends(require_admin)):
    u = await db.users.find_one({"id": user_id}, {"_id": 0})
    if not u:
        raise HTTPException(404, "Användaren hittades inte")
    await db.users.update_one({"id": user_id}, {"$set": {
        "status": "active", "ban_reason": None, "banned_at": None, "banned_by": None,
    }})
    await log_audit(admin, "user_unbanned", {
        "user_id": user_id, "user_email": u.get("email"),
        "old_value": u.get("status"), "new_value": "active",
    })
    return {"ok": True}

@api.delete("/admin/users/{user_id}")
async def delete_user(user_id: str, req: UserDeleteReq, admin: dict = Depends(require_admin)):
    if req.confirmation != "DELETE":
        raise HTTPException(400, "Bekräftelse krävs. Skriv DELETE för att radera permanent.")
    u = await db.users.find_one({"id": user_id}, {"_id": 0})
    if not u:
        raise HTTPException(404, "Användaren hittades inte")
    if u.get("email") == ADMIN_EMAIL:
        raise HTTPException(400, "Det primära administratörskontot kan inte raderas")
    await db.users.delete_one({"id": user_id})
    await db.match_predictions.delete_many({"user_id": user_id})
    await db.tournament_predictions.delete_many({"user_id": user_id})
    await log_audit(admin, "user_deleted", {
        "user_id": user_id, "user_email": u.get("email"),
        "reason": req.reason, "old_value": u.get("status"), "new_value": "deleted",
    })
    return {"ok": True}

@api.post("/admin/invite", dependencies=[Depends(require_admin)])
async def invite_email(req: InviteReq):
    email = req.email.lower()
    await db.invites.update_one({"email": email}, {"$set": {"email": email, "invited_at": now_utc().isoformat()}}, upsert=True)
    return {"ok": True}

@api.get("/admin/invites", dependencies=[Depends(require_admin)])
async def list_invites():
    return await db.invites.find({}, {"_id": 0}).to_list(500)

@api.delete("/admin/invites/{email}", dependencies=[Depends(require_admin)])
async def remove_invite(email: str):
    await db.invites.delete_one({"email": email.lower()})
    return {"ok": True}

@api.post("/admin/strategy-points")
async def set_strategy_points(req: StrategyReq, admin: dict = Depends(require_admin)):
    await db.users.update_one({"id": req.user_id}, {"$set": {"strategy_points": req.points}})
    await log_audit(admin, "strategy_points_override", {"user_id": req.user_id, "points": req.points})
    return {"ok": True}

@api.post("/admin/recompute")
async def admin_recompute(admin: dict = Depends(require_admin)):
    await recompute_live_points()
    await recompute_strategy_points()
    await log_audit(admin, "recompute_all")
    return {"ok": True}

# ============== Admin: Teams (DQ etc.) ==============
@api.post("/admin/teams/disqualify")
async def disqualify_team(req: DisqualifyReq, admin: dict = Depends(require_admin)):
    if not req.reason or not req.reason.strip():
        raise HTTPException(400, "Reason is required")
    t = await db.teams.find_one({"id": req.team_id})
    if not t:
        raise HTTPException(404, "Team not found")
    await db.teams.update_one({"id": req.team_id}, {"$set": {
        "disqualified": True,
        "disqualified_at": now_utc().isoformat(),
        "disqualified_reason": req.reason,
        "disqualified_by": admin.get("name"),
    }})
    await log_audit(admin, "team_disqualify", {"team_id": req.team_id, "team_name": t["team_name"], "reason": req.reason})
    await progress_tournament()
    return {"ok": True}

@api.post("/admin/teams/{team_id}/reinstate")
async def reinstate_team(team_id: str, admin: dict = Depends(require_admin)):
    t = await db.teams.find_one({"id": team_id})
    if not t:
        raise HTTPException(404, "Team not found")
    await db.teams.update_one({"id": team_id}, {"$unset": {"disqualified": "", "disqualified_at": "", "disqualified_reason": "", "disqualified_by": ""}})
    await log_audit(admin, "team_reinstate", {"team_id": team_id, "team_name": t["team_name"]})
    await progress_tournament()
    return {"ok": True}

class ManualAdvanceReq(BaseModel):
    match_id: str
    team_id: str
    side: Literal["home", "away"]
    reason: Optional[str] = None

@api.post("/admin/manual-advance")
async def manual_advance(req: ManualAdvanceReq, admin: dict = Depends(require_admin)):
    m = await db.matches.find_one({"id": req.match_id}, {"_id": 0})
    if not m:
        raise HTTPException(404, "Match not found")
    key = "home_team_id" if req.side == "home" else "away_team_id"
    await db.matches.update_one({"id": req.match_id}, {"$set": {key: req.team_id}})
    await log_audit(admin, "manual_team_assign", {"match_id": req.match_id, "team_id": req.team_id, "side": req.side, "reason": req.reason})
    return {"ok": True}

# ============== Admin: Audit log ==============
@api.get("/admin/audit-log", dependencies=[Depends(require_admin)])
async def get_audit_log(limit: int = 200):
    logs = await db.audit_log.find({}, {"_id": 0}).sort("timestamp", -1).to_list(limit)
    return logs

# ============== Hall of Fame ==============
@api.get("/hall-of-fame")
async def list_hall_of_fame():
    rows = await db.hall_of_fame.find({}, {"_id": 0}).sort("year", -1).to_list(200)
    return rows

@api.post("/hall-of-fame", dependencies=[Depends(require_admin)])
async def create_hall_of_fame(req: HallOfFameReq, admin: dict = Depends(require_admin)):
    existing = await db.hall_of_fame.find_one({"year": req.year})
    if existing:
        raise HTTPException(400, f"Året {req.year} finns redan i Hall of Fame")
    doc = {
        "id": str(uuid.uuid4()),
        "year": req.year,
        "winner_name": req.winner_name,
        "winner_points": req.winner_points,
        "second_name": req.second_name,
        "second_points": req.second_points,
        "third_name": req.third_name,
        "third_points": req.third_points,
        "created_at": now_utc().isoformat(),
    }
    await db.hall_of_fame.insert_one(doc)
    await log_audit(admin, "hall_of_fame_created", {"year": req.year})
    doc.pop("_id", None)
    return doc

@api.delete("/hall-of-fame/{year}", dependencies=[Depends(require_admin)])
async def delete_hall_of_fame(year: int, admin: dict = Depends(require_admin)):
    res = await db.hall_of_fame.delete_one({"year": year})
    if res.deleted_count == 0:
        raise HTTPException(404, "År saknas")
    await log_audit(admin, "hall_of_fame_deleted", {"year": year})
    return {"ok": True}

# ============== Admin: Excel Import ==============
SWE_TZ_OFFSET_HOURS = 2  # Treat Excel times as Europe/Stockholm summer (UTC+2). Store as UTC ISO.

def parse_xlsx(file_bytes: bytes) -> Dict[str, Any]:
    """Parse the Excel file and return preview structure WITHOUT writing to DB."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    teams_data: List[dict] = []
    matches_data: List[dict] = []
    errors: List[str] = []
    if "Teams" in wb.sheetnames:
        ws = wb["Teams"]
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows[1:], 2):
            if not row or not row[0]:
                continue
            try:
                teams_data.append({"team_name": str(row[0]).strip(), "group": str(row[1]).strip().upper(),
                                   "country_code": str(row[2]).strip()})
            except Exception as e:
                errors.append(f"Teams row {i}: {e}")
    if "Matches_Import" in wb.sheetnames:
        ws = wb["Matches_Import"]
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows[1:], 2):
            if not row or row[0] is None:
                continue
            try:
                match_id, stage, round_, group, date, time_, ht, at_, hcc, acc, tv, status, hs, as_, notes = row
                # Date+time -> ISO. Excel stores as date objects or strings.
                if isinstance(date, datetime):
                    d_str = date.strftime("%Y-%m-%d")
                else:
                    d_str = str(date)
                if isinstance(time_, datetime):
                    t_str = time_.strftime("%H:%M")
                else:
                    t_str = str(time_) if time_ else "00:00"
                # Build naive Swedish-local datetime and convert to UTC ISO
                dt_naive = datetime.fromisoformat(f"{d_str}T{t_str}:00")
                dt_utc = dt_naive - timedelta(hours=SWE_TZ_OFFSET_HOURS)
                kickoff_iso = dt_utc.replace(tzinfo=timezone.utc).isoformat()
                m = {
                    "match_number": int(match_id) if match_id is not None else None,
                    "stage": str(stage).strip() if stage else "Group Stage",
                    "round": str(round_).strip() if round_ else None,
                    "group": str(group).strip().upper() if group else None,
                    "kickoff": kickoff_iso,
                    "home_team": str(ht).strip() if ht else None,
                    "away_team": str(at_).strip() if at_ else None,
                    "home_country_code": str(hcc).strip() if hcc else None,
                    "away_country_code": str(acc).strip() if acc else None,
                    "tv_channel": str(tv).strip() if tv else None,
                    "status": str(status).strip() if status else "scheduled",
                    "home_score": hs, "away_score": as_,
                    "notes": str(notes).strip() if notes else None,
                }
                matches_data.append(m)
            except Exception as e:
                errors.append(f"Matches row {i}: {e}")
    # Detect duplicates by (stage, kickoff, home_team, away_team)
    seen = set()
    duplicates = []
    for m in matches_data:
        key = (m["stage"], m["kickoff"], m["home_team"], m["away_team"])
        if key in seen:
            duplicates.append(m["match_number"])
        seen.add(key)
    return {
        "teams": teams_data, "matches": matches_data,
        "groups_detected": sorted({t["group"] for t in teams_data if t.get("group")}),
        "group_matches": sum(1 for m in matches_data if "group" in (m.get("stage") or "").lower()),
        "knockout_matches": sum(1 for m in matches_data if "knockout" in (m.get("stage") or "").lower()),
        "duplicates": duplicates, "errors": errors,
    }

def is_placeholder(label: Optional[str]) -> bool:
    if not label:
        return False
    return bool(PLACEHOLDER_GROUP_RE.match(label) or PLACEHOLDER_THIRD_RE.match(label)
                or PLACEHOLDER_W_RE.match(label) or PLACEHOLDER_RU_RE.match(label))

async def write_import(parsed: Dict[str, Any], replace: bool = True) -> Dict[str, Any]:
    # Build/upsert teams
    if replace:
        await db.teams.delete_many({})
        await db.matches.delete_many({})
        await db.match_predictions.delete_many({})
    team_id_by_name: Dict[str, str] = {}
    for t in parsed["teams"]:
        doc = {"id": str(uuid.uuid4()), "team_name": t["team_name"],
               "country_code": t["country_code"], "group": t["group"]}
        await db.teams.update_one({"team_name": t["team_name"]}, {"$set": doc}, upsert=True)
        existing = await db.teams.find_one({"team_name": t["team_name"]}, {"_id": 0, "id": 1})
        team_id_by_name[t["team_name"]] = existing["id"] if existing else doc["id"]
    inserted = 0
    for m in parsed["matches"]:
        # Determine if home/away are real team names or placeholders
        home_label, away_label = m.get("home_team"), m.get("away_team")
        home_id = team_id_by_name.get(home_label) if home_label and not is_placeholder(home_label) else None
        away_id = team_id_by_name.get(away_label) if away_label and not is_placeholder(away_label) else None
        home_placeholder = home_label if is_placeholder(home_label) else None
        away_placeholder = away_label if is_placeholder(away_label) else None
        stage_lower = (m.get("stage") or "").lower()
        # Map stage label
        if "knockout" in stage_lower or (m.get("round") or "").lower() in ("sextondelsfinal", "åttondelsfinal", "kvartsfinal", "semifinal", "bronsmatch", "final"):
            stage_norm = "Knockout"
        else:
            stage_norm = "Group Stage"
        doc = {
            "id": str(uuid.uuid4()),
            "match_number": m.get("match_number"),
            "stage": stage_norm,
            "round": m.get("round"),
            "group": m.get("group"),
            "kickoff": m.get("kickoff"),
            "home_team_id": home_id, "away_team_id": away_id,
            "home_placeholder": home_placeholder, "away_placeholder": away_placeholder,
            "home_country_code": m.get("home_country_code"),
            "away_country_code": m.get("away_country_code"),
            "tv_channel": m.get("tv_channel"),
            "home_score": m.get("home_score") if isinstance(m.get("home_score"), int) else None,
            "away_score": m.get("away_score") if isinstance(m.get("away_score"), int) else None,
            "status": "finished" if (isinstance(m.get("home_score"), int) and isinstance(m.get("away_score"), int)) else "upcoming",
        }
        await db.matches.insert_one(doc)
        inserted += 1
    return {"teams_inserted": len(parsed["teams"]), "matches_inserted": inserted}

@api.post("/admin/import-preview")
async def import_preview(file: UploadFile = File(...), admin: dict = Depends(require_admin)):
    data = await file.read()
    parsed = parse_xlsx(data)
    # Strip out heavy lists for preview summary
    return {
        "summary": {
            "teams_count": len(parsed["teams"]),
            "matches_count": len(parsed["matches"]),
            "groups_detected": parsed["groups_detected"],
            "group_matches": parsed["group_matches"],
            "knockout_matches": parsed["knockout_matches"],
            "duplicates": parsed["duplicates"],
            "errors": parsed["errors"],
        },
        "teams_sample": parsed["teams"][:6],
        "matches_sample": parsed["matches"][:8],
    }

@api.post("/admin/import-commit")
async def import_commit(file: UploadFile = File(...), replace: bool = Form(True), admin: dict = Depends(require_admin)):
    data = await file.read()
    parsed = parse_xlsx(data)
    if parsed["errors"]:
        return {"ok": False, "errors": parsed["errors"]}
    result = await write_import(parsed, replace=replace)
    await log_audit(admin, "excel_import", {"replace": replace, **result})
    await recompute_live_points()
    await progress_tournament()
    return {"ok": True, **result, "summary": {
        "groups_detected": parsed["groups_detected"],
        "group_matches": parsed["group_matches"],
        "knockout_matches": parsed["knockout_matches"],
        "duplicates": parsed["duplicates"],
    }}

# ============== Seed ==============
async def seed_admin():
    await db.users.create_index("email", unique=True)
    await db.teams.create_index("id", unique=True)
    await db.matches.create_index("id", unique=True)
    await db.match_predictions.create_index([("user_id", 1), ("match_id", 1)], unique=True)
    await db.tournament_predictions.create_index([("user_id", 1), ("version", 1)], unique=True)
    await db.invites.create_index("email", unique=True)
    await db.config.create_index("key", unique=True)
    existing = await db.users.find_one({"email": ADMIN_EMAIL})
    if not existing:
        await db.users.insert_one({
            "id": str(uuid.uuid4()), "name": "Gabriella Bengtsson",
            "email": ADMIN_EMAIL, "password_hash": hash_password(ADMIN_PASSWORD),
            "role": "admin", "created_at": now_utc().isoformat(),
            "live_points": 0, "strategy_points": 0,
        })
        logger.info("Seeded admin user")
    # Default V1 deadline: 5 min before the first WC match (2026-06-11 20:55 Europe/Stockholm = 18:55 UTC)
    cfg = await db.config.find_one({"key": "v1_deadline"})
    if not cfg:
        default_v1 = datetime(2026, 6, 11, 18, 55, 0, tzinfo=timezone.utc).isoformat()
        await db.config.update_one(
            {"key": "v1_deadline"},
            {"$set": {"key": "v1_deadline", "value": default_v1}},
            upsert=True,
        )
        logger.info(f"Seeded default V1 deadline: {default_v1}")

async def seed_from_xlsx_if_needed():
    xlsx_path = DATA_DIR / "vm2026.xlsx"
    if not xlsx_path.exists():
        logger.info("No VM2026 xlsx found; skipping auto-import")
        return
    # If we have fewer than 48 teams or no matches with match_number, force import
    teams_count = await db.teams.count_documents({})
    matches_with_number = await db.matches.count_documents({"match_number": {"$ne": None}})
    if teams_count >= 48 and matches_with_number >= 100:
        logger.info("DB already has VM2026 data; skipping auto-import")
        return
    logger.info("Auto-importing VM2026 fixtures from xlsx...")
    with open(xlsx_path, "rb") as f:
        parsed = parse_xlsx(f.read())
    await write_import(parsed, replace=True)
    await progress_tournament()
    logger.info(f"Auto-imported {len(parsed['teams'])} teams, {len(parsed['matches'])} matches")

@app.on_event("startup")
async def on_startup():
    await seed_admin()
    await seed_from_xlsx_if_needed()

# ============== Mount ==============
app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[
        "http://localhost:3000",
        "http://192.168.0.104:3000",
        "https://bengtssonfamilj.se",
        "https://www.bengtssonfamilj.se",
    ],
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
